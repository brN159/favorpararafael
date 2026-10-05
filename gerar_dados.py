#!/usr/bin/env python3
"""Gera docs/dados.js (usado pelo site) a partir dos PDFs e do Votos.xlsx.
Rode sempre que chegarem boletins novos:  python gerar_dados.py
Precisa estar na mesma pasta do apurar_votos.py."""
import json
from pathlib import Path

from openpyxl import load_workbook

from apurar_votos import CARGOS, NOME_EXCEL, PASTA_PADRAO, carregar_dados

SAIDA = Path(__file__).resolve().parent / "docs" / "dados.js"


def main():
    dados = carregar_dados(PASTA_PADRAO)
    ws = load_workbook(PASTA_PADRAO / NOME_EXCEL, read_only=True)["Votos por Seção"]

    escolas, idx, sec_escola = [], {}, {}
    for row in ws.iter_rows(values_only=True):
        if isinstance(row[0], int):                      # linha de seção
            chave = (row[2], row[3])                     # escola + localidade
            if chave not in idx:
                idx[chave] = len(escolas)
                escolas.append({"nome": row[2], "loc": row[3]})
            sec_escola[row[0]] = idx[chave]

    saida = {"escolas": escolas, "secEscola": sec_escola, "cargos": {}}
    for cargo in CARGOS:
        cands, apurado = {}, {}
        for secao, info in dados.items():
            apurado[secao] = info["resumos"].get(cargo, {}).get("Total Apurado", 0)
            for c_cargo, partido, num, nome, votos in info["candidatos"]:
                if c_cargo == cargo:
                    cands.setdefault(num, [num, nome, partido, {}])[3][secao] = votos
        saida["cargos"][cargo] = {"candidatos": list(cands.values()), "apurado": apurado}

    SAIDA.parent.mkdir(exist_ok=True)
    SAIDA.write_text("window.DADOS=" + json.dumps(saida, ensure_ascii=False, separators=(",", ":")) + ";",
                     encoding="utf-8")
    print(f"{len(dados)} seções -> {SAIDA}")


if __name__ == "__main__":
    main()
