"""Núcleo: baixa/lê o arquivo "Votação por seção eleitoral" do TSE e gera a planilha Excel."""
from __future__ import annotations

import re
import unicodedata
import zipfile
from pathlib import Path
from typing import Callable

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

URL_TSE = "https://cdn.tse.jus.br/estatistica/sead/odsele/votacao_secao/votacao_secao_{ano}_{uf}.zip"
CARGOS = ["Deputado Estadual", "Deputado Federal", "Governador", "Senador", "Presidente"]
REQUIRED = ["NR_TURNO", "NM_MUNICIPIO", "NR_ZONA", "NR_SECAO", "DS_CARGO", "NR_VOTAVEL",
            "NM_VOTAVEL", "QT_VOTOS", "NR_LOCAL_VOTACAO", "NM_LOCAL_VOTACAO"]
OPTIONAL = ["DS_LOCAL_VOTACAO_ENDERECO"]
BASE = ["NR_ZONA", "NR_SECAO", "NR_LOCAL_VOTACAO", "NM_LOCAL_VOTACAO", "DS_LOCAL_VOTACAO_ENDERECO"]


def norm(s) -> str:
    return unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().upper().strip()


def slug(s) -> str:
    return re.sub(r"[^a-z0-9]+", "_", norm(s).lower()).strip("_")


# ---------------------------------------------------------------- download / leitura
def url_zip(ano: int, uf: str) -> str:
    return URL_TSE.format(ano=ano, uf=uf.upper())


def baixar_zip(ano: int, uf: str, pasta, forcar: bool = False,
               progresso: Callable[[float], None] | None = None) -> Path:
    pasta = Path(pasta)
    pasta.mkdir(parents=True, exist_ok=True)
    destino = pasta / f"votacao_secao_{ano}_{uf.upper()}.zip"
    if destino.exists() and not forcar:
        return destino
    url = url_zip(ano, uf)
    tmp = destino.with_suffix(".part")
    with requests.get(url, stream=True, timeout=60, headers={"User-Agent": "Mozilla/5.0"}) as r:
        if r.status_code == 404:
            raise FileNotFoundError(f"O TSE ainda não publicou este arquivo: {url}")
        r.raise_for_status()
        total = int(r.headers.get("content-length", 0))
        lido = 0
        with open(tmp, "wb") as f:
            for bloco in r.iter_content(1 << 20):
                f.write(bloco)
                lido += len(bloco)
                if progresso and total:
                    progresso(min(lido / total, 1.0))
    tmp.replace(destino)
    return destino


def _ler(z: zipfile.ZipFile, nome: str, enc: str, alvo: str, turno: int, progresso):
    total = max(z.getinfo(nome).file_size, 1)
    partes = []
    with z.open(nome) as f:
        leitor = pd.read_csv(f, sep=";", encoding=enc, dtype=str, chunksize=400_000,
                             usecols=lambda c: c in REQUIRED + OPTIONAL)
        primeiro = True
        for ch in leitor:
            if primeiro:
                faltam = [c for c in REQUIRED if c not in ch.columns]
                if faltam:
                    raise ValueError(f"Colunas não encontradas no CSV: {faltam}. "
                                     f"Encontradas: {list(ch.columns)}")
                if "DS_LOCAL_VOTACAO_ENDERECO" not in ch.columns:
                    pass
                primeiro = False
            cargos_ok = {v for v in ch["DS_CARGO"].dropna().unique() if norm(v) == alvo}
            ch = ch[ch["DS_CARGO"].isin(cargos_ok) & (ch["NR_TURNO"].str.strip() == str(turno))]
            if len(ch):
                partes.append(ch.copy())
            if progresso:
                progresso(min(f.tell() / total, 1.0))
    return partes


def carregar(zip_path, cargo: str, turno: int = 1,
             progresso: Callable[[float], None] | None = None) -> pd.DataFrame:
    zip_path = Path(zip_path)
    cache = zip_path.parent / f"{zip_path.stem}_{slug(cargo)}_t{turno}.parquet"
    if cache.exists() and cache.stat().st_mtime >= zip_path.stat().st_mtime:
        return pd.read_parquet(cache)
    alvo = norm(cargo)
    with zipfile.ZipFile(zip_path) as z:
        csvs = [n for n in z.namelist() if n.lower().endswith(".csv")]
        if not csvs:
            raise ValueError("O zip não contém nenhum CSV.")
        nome = max(csvs, key=lambda n: z.getinfo(n).file_size)
        try:
            partes = _ler(z, nome, "utf-8-sig", alvo, turno, progresso)
        except UnicodeDecodeError:
            partes = _ler(z, nome, "latin1", alvo, turno, progresso)
    if not partes:
        raise ValueError(f"Nenhuma linha para o cargo '{cargo}' no turno {turno}.")
    df = pd.concat(partes, ignore_index=True)
    if "DS_LOCAL_VOTACAO_ENDERECO" not in df.columns:
        df["DS_LOCAL_VOTACAO_ENDERECO"] = ""
    df["DS_LOCAL_VOTACAO_ENDERECO"] = df["DS_LOCAL_VOTACAO_ENDERECO"].fillna("")
    for c in ["NR_ZONA", "NR_SECAO", "QT_VOTOS"]:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0).astype(int)
    df["NR_LOCAL_VOTACAO"] = pd.to_numeric(df["NR_LOCAL_VOTACAO"], errors="coerce").fillna(0).astype(int)
    df["NM_VOTAVEL"] = df["NM_VOTAVEL"].fillna("")
    df = df.drop(columns=["NR_TURNO"])
    df.to_parquet(cache)
    return df


# ---------------------------------------------------------------- consultas
def municipios(df: pd.DataFrame) -> list[str]:
    return sorted(df["NM_MUNICIPIO"].dropna().unique())


def rotulo(nome: str, numero) -> str:
    return f"{str(nome).title()} ({numero})"


def candidatos(df: pd.DataFrame, municipio: str, incluir_brancos_nulos: bool = False) -> pd.DataFrame:
    sub = df[df["NM_MUNICIPIO"] == municipio]
    g = sub.groupby(["NR_VOTAVEL", "NM_VOTAVEL"], as_index=False)["QT_VOTOS"].sum()
    if not incluir_brancos_nulos:
        g = g[~g["NM_VOTAVEL"].map(lambda s: "BRANCO" in norm(s) or "NULO" in norm(s))]
    g["rotulo"] = [rotulo(n, r) for r, n in zip(g["NR_VOTAVEL"], g["NM_VOTAVEL"])]
    return g.sort_values("QT_VOTOS", ascending=False).reset_index(drop=True)


def completar_locais(df: pd.DataFrame, pasta=".") -> pd.DataFrame:
    """Preenche nome/endereço do local quando o TSE manda '#NULO#', usando locais_*.csv (por zona e seção)."""
    ruim = lambda s: s.isna() | s.astype(str).str.strip().isin(["", "#NULO#", "#NE#", "#NULO", "nan"])
    df = df.copy()
    chaves = ["NM_MUNICIPIO", "NR_ZONA", "NR_SECAO"]
    arquivos = sorted(Path(pasta).glob("locais_*.csv"))
    if arquivos:
        dp = pd.concat([pd.read_csv(a, sep=";", dtype=str, encoding="utf-8-sig") for a in arquivos])
        dp["NM_MUNICIPIO"] = dp["NM_MUNICIPIO"].str.strip()
        dp["NR_ZONA"] = pd.to_numeric(dp["NR_ZONA"]).astype(int)
        dp["NR_SECAO"] = pd.to_numeric(dp["NR_SECAO"]).astype(int)
        dp = dp.rename(columns={"NM_LOCAL_VOTACAO": "_nome", "DS_LOCAL_VOTACAO_ENDERECO": "_end"})
        dp = dp[chaves + ["_nome", "_end"]].drop_duplicates(chaves)
        df = df.merge(dp, on=chaves, how="left")
        m = ruim(df["NM_LOCAL_VOTACAO"]) & df["_nome"].notna()
        df.loc[m, "NM_LOCAL_VOTACAO"] = df.loc[m, "_nome"]
        m = ruim(df["DS_LOCAL_VOTACAO_ENDERECO"]) & df["_end"].notna()
        df.loc[m, "DS_LOCAL_VOTACAO_ENDERECO"] = df.loc[m, "_end"]
        df = df.drop(columns=["_nome", "_end"])
    m = ruim(df["NM_LOCAL_VOTACAO"])
    df.loc[m, "NM_LOCAL_VOTACAO"] = "Local " + df.loc[m, "NR_LOCAL_VOTACAO"].astype(str)
    m = ruim(df["DS_LOCAL_VOTACAO_ENDERECO"])
    df.loc[m, "DS_LOCAL_VOTACAO_ENDERECO"] = ""
    return df


def tabela_secoes(df: pd.DataFrame, municipio: str, numeros: list[str]):
    sub = df[df["NM_MUNICIPIO"] == municipio]
    base = (sub[BASE].drop_duplicates(["NR_ZONA", "NR_SECAO"])
            .sort_values(["NR_ZONA", "NR_SECAO"]).reset_index(drop=True))
    sel = sub[sub["NR_VOTAVEL"].isin([str(n) for n in numeros])].copy()
    sel["rotulo"] = [rotulo(n, r) for r, n in zip(sel["NR_VOTAVEL"], sel["NM_VOTAVEL"])]
    ordem = []
    for n in numeros:
        r = sel.loc[sel["NR_VOTAVEL"] == str(n), "rotulo"]
        if len(r):
            ordem.append(r.iloc[0])
    if sel.empty:
        return base, []
    piv = sel.pivot_table(index=["NR_ZONA", "NR_SECAO"], columns="rotulo", values="QT_VOTOS",
                          aggfunc="sum", fill_value=0).reset_index()
    t = base.merge(piv, on=["NR_ZONA", "NR_SECAO"], how="left")
    t[ordem] = t[ordem].fillna(0).astype(int)
    return t[BASE + ordem], ordem


def tabela_escolas(t: pd.DataFrame, rotulos: list[str]) -> pd.DataFrame:
    agg = {"NM_LOCAL_VOTACAO": "first", "DS_LOCAL_VOTACAO_ENDERECO": "first", "NR_SECAO": "count"}
    agg.update({r: "sum" for r in rotulos})
    e = t.groupby(["NR_ZONA", "NR_LOCAL_VOTACAO"], as_index=False).agg(agg)
    e = e.rename(columns={"NR_SECAO": "N_SECOES"})
    if rotulos:
        e = e.sort_values(rotulos[0], ascending=False)
    return e.reset_index(drop=True)


# ---------------------------------------------------------------- Excel
_F = "Arial"
_HF = PatternFill("solid", fgColor="1F4E78")
_thin = Side(style="thin", color="BFBFBF")
_BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _hdr(ws, row, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row, i, c)
        x.font = Font(name=_F, bold=True, color="FFFFFF")
        x.fill = _HF
        x.border = _BD
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def gerar_excel(t: pd.DataFrame, rotulos: list[str], titulo: str) -> bytes:
    import io
    wb = Workbook()
    we = wb.active
    we.title = "Por Escola"
    ws = wb.create_sheet("Por Seção")
    n, r0 = len(t), 5
    rl = r0 + n - 1
    k = len(rotulos)
    c0 = 6  # primeira coluna de votos (F) nas duas abas

    # ---- Por Seção
    ws["A1"] = titulo
    ws["A1"].font = Font(name=_F, bold=True, size=13)
    _hdr(ws, 4, ["Zona", "Seção", "Cód. local", "Escola / Local de votação", "Endereço"] + rotulos)
    for i, row in enumerate(t.values.tolist()):
        for j, v in enumerate(row, 1):
            x = ws.cell(r0 + i, j, v)
            x.font = Font(name=_F)
            x.border = _BD
            if j >= c0:
                x.number_format = "#,##0"
            if j <= 3:
                x.alignment = Alignment(horizontal="center")
    rt = rl + 1
    ws.cell(rt, 4, "TOTAL").font = Font(name=_F, bold=True)
    for j in range(c0, c0 + k):
        L = get_column_letter(j)
        x = ws.cell(rt, j, f"=SUM({L}{r0}:{L}{rl})")
        x.font = Font(name=_F, bold=True)
        x.number_format = "#,##0"
    for col, w in zip("ABCDE", [7, 8, 11, 55, 40]):
        ws.column_dimensions[col].width = w
    for j in range(c0, c0 + k):
        ws.column_dimensions[get_column_letter(j)].width = 20
    ws.freeze_panes = "A5"

    # ---- Por Escola
    e = tabela_escolas(t, rotulos)
    ne = len(e)
    we["A1"] = titulo
    we["A1"].font = Font(name=_F, bold=True, size=13)
    we["A2"] = "Votos somados por fórmula a partir da aba 'Por Seção'."
    we["A2"].font = Font(name=_F, italic=True, size=9)
    _hdr(we, 4, ["Zona", "Cód. local", "Escola / Local de votação", "Endereço", "Nº seções"]
         + rotulos + [f"% {r}" for r in rotulos])
    S = "'Por Seção'!"
    rng = lambda col: f"{S}${col}${r0}:${col}${rl}"
    tr = 5 + ne
    for i, row in e.iterrows():
        r = 5 + i
        vals = [int(row["NR_ZONA"]), int(row["NR_LOCAL_VOTACAO"]), row["NM_LOCAL_VOTACAO"],
                row["DS_LOCAL_VOTACAO_ENDERECO"]]
        for j, v in enumerate(vals, 1):
            we.cell(r, j, v)
        we.cell(r, 5, f"=COUNTIFS({rng('A')},A{r},{rng('C')},B{r})")
        for q in range(k):
            L = get_column_letter(c0 + q)
            we.cell(r, c0 + q, f"=SUMIFS({S}{L}${r0}:{L}${rl},{rng('A')},$A{r},{rng('C')},$B{r})")
            we.cell(r, c0 + k + q, f"=IF({L}${tr}=0,0,{L}{r}/{L}${tr})")
        for j in range(1, c0 + 2 * k):
            x = we.cell(r, j)
            x.font = Font(name=_F)
            x.border = _BD
            if c0 <= j < c0 + k:
                x.number_format = "#,##0"
            elif j >= c0 + k:
                x.number_format = "0.0%"
            elif j in (1, 2, 5):
                x.alignment = Alignment(horizontal="center")
    we.cell(tr, 3, "TOTAL")
    we.cell(tr, 5, f"=SUM(E5:E{tr - 1})")
    for q in range(k):
        L = get_column_letter(c0 + q)
        P = get_column_letter(c0 + k + q)
        we.cell(tr, c0 + q, f"=SUM({L}5:{L}{tr - 1})").number_format = "#,##0"
        we.cell(tr, c0 + k + q, f"=SUM({P}5:{P}{tr - 1})").number_format = "0.0%"
    for j in range(1, c0 + 2 * k):
        we.cell(tr, j).font = Font(name=_F, bold=True)
        we.cell(tr, j).border = _BD
    if k:
        a, b = get_column_letter(c0), get_column_letter(c0 + k - 1)
        we.cell(tr + 1, 3, "Conferência (Por Escola − Por Seção; deve ser 0):").font = Font(name=_F, italic=True, size=9)
        we.cell(tr + 1, 6, f"=SUM({a}{tr}:{b}{tr})-SUM({S}{a}{rt}:{b}{rt})").font = Font(name=_F, size=9)
    for col, w in zip("ABCDE", [7, 11, 55, 40, 10]):
        we.column_dimensions[col].width = w
    for j in range(c0, c0 + 2 * k):
        we.column_dimensions[get_column_letter(j)].width = 20 if j < c0 + k else 14
    we.freeze_panes = "A5"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def gerar_por_candidato(t: pd.DataFrame, rotulos: list[str], titulo_base: str) -> dict[str, bytes]:
    out = {}
    for r in rotulos:
        out[f"{slug(r)}.xlsx"] = gerar_excel(t[BASE + [r]], [r], f"{r} – {titulo_base}")
    return out
